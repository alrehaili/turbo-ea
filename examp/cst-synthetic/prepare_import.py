"""Validate an editable collection workbook and package it for Workspace transfer.

Usage: python prepare_import.py collection.xlsx output.zip [--check-only]
Requires openpyxl. No network access, database access or application credentials.
"""
import argparse
import io
import json
import math
import sys
import uuid
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

NAMESPACE=uuid.UUID('d5bc7f6e-333c-5ebf-9056-40d83e4a1f61')
def stable(s):return str(uuid.uuid5(NAMESPACE,s))
def datevalue(v):
    if v is None or v=='':return None
    if isinstance(v,(date,datetime)):return v.date().isoformat() if isinstance(v,datetime) else v.isoformat()
    return date.fromisoformat(str(v).strip()[:10]).isoformat()
def jsonvalue(v,default):
    if v is None or v=='':return default
    return v if isinstance(v,(dict,list)) else json.loads(str(v))
def records(ws,header_row):
    rows=ws.iter_rows(min_row=header_row,values_only=True)
    keys=[str(x or '').strip() for x in next(rows,[])]
    for rn,row in enumerate(rows,header_row+1):
        if any(v is not None and v!='' for v in row):yield rn,{keys[i]:v for i,v in enumerate(row) if i<len(keys) and keys[i]}
def value(v,f):
    if v is None or v=='':
        if f.get('required') and not f.get('readonly'):raise ValueError(f"Required field {f['key']} is empty")
        return None
    typ=f['type']
    if typ in ('number','cost'):
        if isinstance(v,bool):raise ValueError('A number cannot be TRUE/FALSE')
        n=float(v)
        if not math.isfinite(n):raise ValueError('Number must be finite')
        return int(n) if n.is_integer() else n
    if typ=='boolean':
        if isinstance(v,bool):return v
        if str(v).lower() in ('true','1','yes'):return True
        if str(v).lower() in ('false','0','no'):return False
        raise ValueError('Boolean must be TRUE or FALSE')
    if typ=='date':return datevalue(v)
    if typ in ('single_select','multiple_select'):
        valid={str(o['key']) for o in f.get('options',[])}
        result=jsonvalue(v,[]) if typ=='multiple_select' else str(v)
        vs=result if typ=='multiple_select' else [result]
        if not isinstance(vs,list) or any(str(x) not in valid for x in vs):raise ValueError(f"Use stored option keys for {f['key']}")
        return result
    return str(v)

def collect(source,schema):
    wb=load_workbook(source,read_only=True,data_only=True)
    cards=[];owners=[];errors=[];ids=set();refs={}
    types={t['key']:t for t in schema['types']}
    for tk,t in types.items():
        name=tk[:31]
        if name not in wb.sheetnames:errors.append(f'Missing type sheet: {name}');continue
        fs=[f for s in t.get('fields_schema') or [] for f in s.get('fields',[])]
        for rn,row in records(wb[name],4):
            if not row.get('name'):
                errors.append(f'{name} row {rn}: populated row has no card name');continue
            try:
                ext=str(row.get('external_id') or '').strip()
                if not ext:raise ValueError('external_id is required')
                if ext in ids:raise ValueError('duplicate external_id')
                ids.add(ext)
                subtype=row.get('subtype') or None
                if subtype and subtype not in {s['key'] for s in t.get('subtypes') or []}:raise ValueError('unknown subtype key')
                attrs={}
                for f in fs:
                    v=value(row.get(f['key']),f)
                    if v is not None:attrs[f['key']]=v
                parent=str(row.get('parent_path') or '').strip()
                if any(x in str(row['name']) for x in ('\\',' / ')):raise ValueError('Use a card name without backslashes or the reserved path separator " / "')
                state=row.get('architecture_state') or 'current'
                if state not in ('current','transition','target'):raise ValueError('unknown architecture_state')
                c={'type':tk,'name':str(row['name']).strip(),'parent_path':parent,'subtype':subtype,'description':str(row.get('description') or ''),'external_id':ext,'alias':None,'approval_status':'DRAFT','status':'ACTIVE','source_system':'Organization collection workbook','confidence':'medium','architecture_state':state,'change_type':None,'lifecycle':{},'attributes':attrs,'created_via':'api','source_id':stable('collection:'+ext)}
                ownref=(parent+' / ' if parent else '')+c['name']
                if (tk,ownref) in refs:raise ValueError('duplicate type/path/name')
                refs[(tk,ownref)]=c
                cards.append(c)
                if row.get('owner_email'):owners.append((c,str(row['owner_email']).strip().lower()))
            except (ValueError,TypeError,json.JSONDecodeError) as exc:errors.append(f'{name} row {rn}: {exc}')
    byid={c['external_id']:c for c in cards}
    for sheet in ('Lifecycle','Card Details'):
        if sheet not in wb.sheetnames:continue
        for rn,row in records(wb[sheet],1):
            try:
                c=byid.get(str(row.get('external_id') or ''))
                if c is None:raise ValueError('external_id does not match a type-sheet record')
                if sheet=='Lifecycle':c['lifecycle']={k:v for k in ('plan','phaseIn','active','phaseOut','endOfLife') if (v:=datevalue(row.get(k))) is not None}
                else:
                    for k in ('alias','approval_status','status','change_type','source_system','confidence'):
                        if row.get(k) is not None and row[k]!='':c[k]=row[k]
                    if c['approval_status'] not in ('DRAFT','IN_REVIEW','APPROVED','REJECTED'):raise ValueError('unknown approval_status')
                    if c['status'] not in ('ACTIVE','ARCHIVED'):raise ValueError('unknown status')
            except (ValueError,TypeError) as exc:errors.append(f'{sheet} row {rn}: {exc}')
    for c in cards:
        if c['parent_path'] and (c['type'],c['parent_path']) not in refs:errors.append(f"Parent not included: {c['parent_path']}")
        if 'hierarchyLevel' in c['attributes']:c['attributes']['hierarchyLevel']=len(c['parent_path'].split(' / '))+1 if c['parent_path'] else 1
    reltypes={r['key']:r for r in schema['relation_types']};rels=[];seen=set()
    if 'Relations' in wb.sheetnames:
        for rn,row in records(wb['Relations'],1):
            try:
                rt=reltypes.get(str(row.get('type') or ''))
                if rt is None:raise ValueError('unknown relation key')
                sk=(str(row.get('source_type') or ''),str(row.get('source_ref') or ''))
                tk=(str(row.get('target_type') or ''),str(row.get('target_ref') or ''))
                if sk[0]!=rt['source_type_key'] or tk[0]!=rt['target_type_key']:raise ValueError('endpoint types do not match the relation type')
                if sk not in refs or tk not in refs:raise ValueError('endpoint reference is not included in the collection workbook')
                source,target=refs[sk],refs[tk]
                key=(rt['key'],source['source_id'],target['source_id'])
                if key in seen:raise ValueError('duplicate relation')
                seen.add(key)
                attrs=jsonvalue(row.get('attributes'),{})
                if not isinstance(attrs,dict):raise ValueError('attributes must be a JSON object')
                for f in rt.get('attributes_schema') or []:
                    if f['key'] in attrs:attrs[f['key']]=value(attrs[f['key']],f)
                rels.append({'type':rt['key'],'source_type':sk[0],'source_ref':sk[1],'target_type':tk[0],'target_ref':tk[1],'description':row.get('description') or '', 'attributes':attrs,'source_id':stable('collection-rel:'+':'.join(key)),'source_card_id':source['source_id'],'target_card_id':target['source_id']})
            except (ValueError,TypeError,json.JSONDecodeError) as exc:errors.append(f'Relations row {rn}: {exc}')
    stakeholders=[]
    for c,email in owners:
        if '@' not in email:errors.append('Invalid owner email on '+c['external_id']);continue
        rs=types[c['type']].get('stakeholder_roles') or []
        role=next((r['key'] for r in rs if r['key']=='responsible'),rs[0]['key'] if rs else 'responsible')
        stakeholders.append({'id':stable('collection-owner:'+c['external_id']),'role':role,'card_id__ref':(c['parent_path']+' / ' if c['parent_path'] else '')+c['name'],'card_id__type':c['type'],'card_id__id':c['source_id'],'user_id__email':email})
    wb.close()
    if errors:raise ValueError('\n'.join(errors[:50]))
    return cards,rels,stakeholders

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path);parser.add_argument('output',type=Path);parser.add_argument('--check-only',action='store_true')
    args=parser.parse_args()
    schema=json.loads(Path(__file__).with_name('collection-schema.json').read_text(encoding='utf-8'))
    try:cards,rels,stakeholders=collect(args.source,schema)
    except ValueError as exc:print(str(exc),file=sys.stderr);return 1
    print(f'Validated: {len(cards)} cards, {len(rels)} relations, {len(stakeholders)} owner references. No users are created by this converter.')
    if args.check_only:return 0
    if not cards:print('No records to import. Fill at least one type-sheet row.',file=sys.stderr);return 1
    sheets={'Cards':(schema['card_columns'],cards),'Relations':(schema['relation_columns'],rels),'Stakeholders':(['id','role','card_id__ref','card_id__type','card_id__id','user_id__email'],stakeholders)}
    wb=Workbook();wb.remove(wb.active)
    for name,(cols,rows) in sheets.items():
        ws=wb.create_sheet(name);ws.append(cols);ws.freeze_panes='A2'
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17324D')
        for row in rows:ws.append([json.dumps(row[c],ensure_ascii=False) if isinstance(row.get(c),(dict,list)) else row.get(c) for c in cols])
    stream=io.BytesIO();wb.save(stream)
    manifest={'format_version':'1','app_version':'2.43.0','exported_at':datetime.now(timezone.utc).isoformat(),'source_url':'collection-workbook','sections':list(sheets),'notes':'Cards, relationships and existing-user owner references only. No settings, metamodel or account changes.'}
    args.output.parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(args.output,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('manifest.json',json.dumps(manifest,indent=2));z.writestr('workspace.xlsx',stream.getvalue())
    print('Created '+str(args.output))
    return 0
if __name__=='__main__':sys.exit(main())
